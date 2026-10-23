"""Adapt-material flow. Gemini is replaced by fakes: no network, no API usage."""
import io
import json
from datetime import date
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document
from PIL import Image

from apps.accounts.models import Teacher
from apps.materials.models import MaterialStatus, OriginalMaterial, Subject
from apps.students.models import Student

from . import ai
from .models import AdaptationStatus, CurricularAdaptation

PASSWORD = 'Aula-Inclusiva-2026'

ANALYSIS = {
    'readable': True, 'problem': '', 'extracted_text': 'Juan tiene 3/4 de pizza y regala 2/4.',
    'figures': [{'description': 'Círculo dividido en 4 partes, 3 pintadas.', 'source': 1, 'box_2d': [100, 100, 900, 900]}],
    'exercise_count': 1,
    'detected_barriers': ['Texto denso'], 'cognitive_load': 'ALTA', 'summary': 'Resta de fracciones.',
}
ADAPTED = {
    'title': 'La pizza de Juan', 'subtitle': '', 'summary': 'Fracciones con apoyo visual.', 'intro': 'Vamos paso a paso.',
    'learning_goal': 'Vas a aprender a restar fracciones.',
    'concepts': [
        {'title': '¿Qué es una fracción?', 'explanation': 'Una fracción son partes iguales de un entero.',
         'example': 'Como repartir una pizza.', 'figure_number': 1},
        {'title': 'Restar', 'explanation': 'Restar es quitar.', 'example': 'Si comes 2 porciones, quedan menos.',
         'figure_number': 1},
    ],
    'check_questions': [
        {'question': '¿Cuántas partes tiene la pizza?', 'options': ['4', '3', '2'], 'answer': '4'},
        {'question': 'Pregunta sin opciones', 'options': ['solo una'], 'answer': 'x'},
    ],
    'steps': [
        {'title': 'Paso 1: Mira la pizza', 'instruction': 'Cuenta las porciones.', 'answer_type': 'ESPACIO',
         'keyword': 'porciones', 'visual_support': 'Círculo en 4 partes.', 'figure_description': '', 'figure_number': 0,
         'hint': 'Recuerda qué es una fracción.'},
        {'title': 'Resta', 'instruction': 'Quita 2 porciones.', 'answer_type': 'NINGUNA', 'figure_number': 0,
         'visual': {'type': 'FRACTION_CIRCLE', 'parts': 4, 'shaded': 3, 'crossed': 2}},
        {'title': 'Inventado', 'instruction': 'x', 'answer_type': 'NINGUNA', 'figure_number': 9,
         'visual': {'type': 'MAGIC'}},
    ],
    'glossary': [], 'closing_message': '¡Muy bien!', 'udl_principle': 'Representación',
    'udl_justification': 'Se ancla la fracción a un objeto concreto.',
}


def fake_analyze(**kwargs):
    return dict(ANALYSIS), 'gemini-test'


def fake_adapt(**kwargs):
    figure_count = len(kwargs['analysis'].get('figures', []))
    return ai.normalize_adaptation(json.loads(json.dumps(ADAPTED)), figure_count=figure_count), 'gemini-test'


def png_bytes(size=200):
    buffer = io.BytesIO()
    Image.new('RGB', (size, size), 'white').save(buffer, 'PNG')
    return buffer.getvalue()


@override_settings(AI_PROCESS_INLINE=True)
@mock.patch('apps.adaptations.services.ai.adapt_worksheet', side_effect=fake_adapt)
@mock.patch('apps.adaptations.services.ai.analyze_material', side_effect=fake_analyze)
class AdaptFlowTests(TestCase):

    def setUp(self):
        self.teacher = Teacher.objects.create_user('ana@correo.com', PASSWORD, full_name='Ana Pérez')
        self.client.force_login(self.teacher)
        self.student = Student.objects.create(teacher=self.teacher, name='Mateo Benítez', birth_date=date(2016, 5, 1),
                                              grade='EGB_4', conditions=['DYSLEXIA', 'DYSCALCULIA'])
        self.math = Subject.objects.get(code='MAT')

    def form_data(self, **changes):
        data = {'students': [str(self.student.pk)], 'subject': self.math.pk, 'title': 'Fracciones',
                'topic': 'Fracciones simples', 'input_type': 'TEXT', 'input_text': 'Juan tiene 3/4 de pizza...',
                f'condition_{self.student.pk}': 'DYSCALCULIA'}
        data.update(changes)
        return data

    def test_text_material_creates_adaptation(self, analyze, adapt):
        response = self.client.post(reverse('adapt'), self.form_data())
        material = OriginalMaterial.objects.get()
        self.assertRedirects(response, reverse('material_detail', args=[material.pk]))
        self.assertEqual(material.status, MaterialStatus.PROCESSED)
        adaptation = CurricularAdaptation.objects.get()
        self.assertEqual(adaptation.condition, 'DYSCALCULIA')
        self.assertEqual(adaptation.version, 1)
        self.assertEqual(adaptation.status, AdaptationStatus.IN_REVIEW)
        self.assertRegex(adaptation.code, r'^ADAP-MAT-')
        self.assertEqual(adaptation.content['steps'][0]['title'], 'Mira la pizza')  # "Paso 1:" removed

    def test_student_name_never_reaches_the_ai(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        sent = json.dumps([analyze.call_args.kwargs, adapt.call_args.kwargs], default=str)
        self.assertNotIn('Mateo', sent)
        self.assertNotIn('Benítez', sent)

    def test_recommended_adjustments_when_teacher_picks_none(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        self.assertIn('ACCESSIBLE_FONT', CurricularAdaptation.objects.get().applied_adjustments)

    def test_photo_is_sent_as_image_and_figure_is_cropped(self, analyze, adapt):
        photo = SimpleUploadedFile('guia.png', png_bytes(), content_type='image/png')
        self.client.post(reverse('adapt'), self.form_data(input_type='PHOTO', input_text='', file=photo))
        parts = analyze.call_args.kwargs['image_parts']
        self.assertEqual(parts[1]['inlineData']['mimeType'], 'image/jpeg')
        material = OriginalMaterial.objects.get()
        self.assertEqual(material.file_name, 'guia.png')
        self.assertTrue(material.ai_analysis['figures'][0]['image'].endswith('.png'))
        figure = self.client.get(reverse('material_figure', args=[material.pk, 1]))
        self.assertEqual(figure['Content-Type'], 'image/png')
        page = self.client.get(reverse('material_detail', args=[material.pk]))
        self.assertContains(page, 'Figura de la guía original')
        self.assertContains(page, 'Círculo de fracciones dibujado por el sistema')
        for section in ('Vas a aprender a restar fracciones.', 'Aprende', 'Ahora practica',
                        'Comprueba lo que aprendiste', 'Pista:', 'Como repartir una pizza.'):
            self.assertContains(page, section)

    def test_pdf_pages_are_sent_as_images(self, analyze, adapt):
        buffer = io.BytesIO()
        Image.new('RGB', (200, 280), 'white').save(buffer, 'PDF')
        pdf = SimpleUploadedFile('guia.pdf', buffer.getvalue(), content_type='application/pdf')
        self.client.post(reverse('adapt'), self.form_data(input_type='PDF', input_text='', file=pdf))
        self.assertEqual(len(analyze.call_args.kwargs['image_parts']), 2)  # "Imagen 1:" + the page
        self.assertEqual(OriginalMaterial.objects.get().status, MaterialStatus.PROCESSED)

    def test_invalid_figure_numbers_and_visuals_are_dropped(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())  # text input: no figures to show
        content = CurricularAdaptation.objects.get().content
        steps = content['steps']
        self.assertEqual(content['concepts'][0]['figure_number'], 1)
        self.assertEqual(content['concepts'][1]['figure_number'], 0)  # each figure is used only once
        self.assertEqual(len(content['check_questions']), 1)  # a question needs at least 2 options
        self.assertEqual(steps[0]['figure_number'], 0)
        self.assertEqual(steps[1]['visual'], {'type': 'FRACTION_CIRCLE', 'parts': 4, 'shaded': 3, 'crossed': 2})
        self.assertEqual(steps[2]['figure_number'], 0)
        self.assertIsNone(steps[2]['visual'])

    def test_word_file_is_read_as_text(self, analyze, adapt):
        document = Document()
        document.add_paragraph('Resuelve: 3/4 - 2/4')
        document.add_picture(io.BytesIO(png_bytes(size=80)))
        buffer = io.BytesIO()
        document.save(buffer)
        word = SimpleUploadedFile('guia.docx', buffer.getvalue())
        self.client.post(reverse('adapt'), self.form_data(input_type='WORD', input_text='', file=word))
        self.assertIn('3/4 - 2/4', analyze.call_args.kwargs['text'])
        self.assertEqual(len(analyze.call_args.kwargs['image_parts']), 2)  # the embedded picture

    def test_file_must_match_input_type(self, analyze, adapt):
        fake_pdf = SimpleUploadedFile('guia.png', png_bytes())
        response = self.client.post(reverse('adapt'), self.form_data(input_type='PDF', input_text='', file=fake_pdf))
        self.assertContains(response, 'Este tipo de entrada acepta archivos PDF')
        self.assertFalse(OriginalMaterial.objects.exists())

    def test_ai_failure_marks_material_as_error(self, analyze, adapt):
        analyze.side_effect = ai.AIError('La foto está borrosa.')
        self.client.post(reverse('adapt'), self.form_data())
        material = OriginalMaterial.objects.get()
        self.assertEqual(material.status, MaterialStatus.ERROR)
        self.assertEqual(material.ai_error, 'La foto está borrosa.')
        page = self.client.get(reverse('material_detail', args=[material.pk]))
        self.assertContains(page, 'La foto está borrosa.')

    def test_status_endpoint(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        material = OriginalMaterial.objects.get()
        data = self.client.get(reverse('material_status', args=[material.pk])).json()
        self.assertEqual(data['status'], 'PROCESSED')

    def test_pdf_requires_approval(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        adaptation = CurricularAdaptation.objects.get()
        blocked = self.client.get(reverse('adaptation_pdf', args=[adaptation.pk]))
        self.assertEqual(blocked.status_code, 302)
        self.client.post(reverse('adaptation_action', args=[adaptation.pk]), {'action': 'approve'})
        response = self.client.get(reverse('adaptation_pdf', args=[adaptation.pk]))
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        adaptation.refresh_from_db()
        self.assertEqual(adaptation.status, AdaptationStatus.EXPORTED)

    def test_regenerate_creates_new_version_without_the_name(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        first = CurricularAdaptation.objects.get()
        self.client.post(reverse('adaptation_action', args=[first.pk]),
                         {'action': 'regenerate', 'feedback': 'Mateo se confundió con los cuartos'})
        second = CurricularAdaptation.objects.get(version=2)
        self.assertEqual(second.source_adaptation, first)
        self.assertTrue(adapt.call_args.kwargs['simplify'])
        self.assertNotIn('Mateo', adapt.call_args.kwargs['feedback'])

    def test_color_filter_and_manual_edit(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        adaptation = CurricularAdaptation.objects.get()
        self.client.post(reverse('adaptation_action', args=[adaptation.pk]),
                         {'action': 'color_filter', 'color_filter': 'CREAM'})
        self.client.post(reverse('adaptation_edit', args=[adaptation.pk]),
                         {'title': 'Nuevo título', 'intro': 'Hola', 'step_0_title': 'Cuenta', 'step_0_instruction': 'Cuenta 3.',
                          'concept_0_explanation': 'Partes iguales.', 'concept_1_delete': 'on'})
        adaptation.refresh_from_db()
        self.assertEqual(adaptation.color_filter, 'CREAM')
        self.assertEqual(adaptation.title, 'Nuevo título')
        self.assertEqual(adaptation.content['steps'][0]['instruction'], 'Cuenta 3.')
        self.assertEqual(adaptation.content['concepts'][0]['explanation'], 'Partes iguales.')
        self.assertEqual(len(adaptation.content['concepts']), 1)
        self.assertEqual(adaptation.status, AdaptationStatus.IN_REVIEW)

    def test_other_teacher_cannot_open_material(self, analyze, adapt):
        self.client.post(reverse('adapt'), self.form_data())
        material = OriginalMaterial.objects.get()
        adaptation = CurricularAdaptation.objects.get()
        intruder = Teacher.objects.create_user('luis@correo.com', PASSWORD, full_name='Luis Mora')
        self.client.force_login(intruder)
        for url in (reverse('material_detail', args=[material.pk]), reverse('material_status', args=[material.pk]),
                    reverse('material_figure', args=[material.pk, 1]), reverse('adaptation_pdf', args=[adaptation.pk])):
            self.assertEqual(self.client.get(url).status_code, 404)

    def test_cannot_adapt_for_another_teachers_student(self, analyze, adapt):
        other = Teacher.objects.create_user('luis@correo.com', PASSWORD, full_name='Luis Mora')
        foreign = Student.objects.create(teacher=other, name='Daniel Vera', birth_date=date(2016, 1, 1),
                                         grade='EGB_4', conditions=['ADHD'])
        response = self.client.post(reverse('adapt'), self.form_data(students=[str(foreign.pk)]))
        self.assertEqual(response.status_code, 200)  # form error, nothing created
        self.assertFalse(OriginalMaterial.objects.exists())


class AIClientTests(TestCase):

    @override_settings(GEMINI_API_KEY='')
    def test_missing_key_is_reported(self):
        with self.assertRaisesMessage(ai.AIError, 'GEMINI_API_KEY'):
            ai.generate_json('system', [{'text': 'hola'}], ai.ANALYSIS_SCHEMA)

    @override_settings(GEMINI_API_KEY='test-key')
    def test_unreadable_photo_raises_error(self):
        with mock.patch('apps.adaptations.ai._call', return_value={
            'candidates': [{'content': {'parts': [{'text': json.dumps({**ANALYSIS, 'readable': False, 'problem': 'Borrosa'})}]}}],
        }):
            with self.assertRaisesMessage(ai.AIError, 'Borrosa'):
                ai.analyze_material(text='x', subject='Matemática', topic='Fracciones')

    @override_settings(GEMINI_API_KEY='test-key')
    def test_falls_back_to_second_model_when_overloaded(self):
        import urllib.error
        overloaded = urllib.error.HTTPError('url', 503, 'busy', {}, None)
        ok = {'candidates': [{'content': {'parts': [{'text': '{"readable": true}'}]}}], 'modelVersion': 'lite'}
        with mock.patch('apps.adaptations.ai._call', side_effect=[overloaded, ok]) as call:
            data, model = ai.generate_json('system', [{'text': 'hola'}], ai.ANALYSIS_SCHEMA)
        self.assertEqual(model, 'lite')
        self.assertEqual(call.call_count, 2)


class VisualsTests(TestCase):

    def test_clean_visual_clamps_and_rejects(self):
        from . import visuals
        self.assertIsNone(visuals.clean_visual({'type': 'NONE'}))
        self.assertIsNone(visuals.clean_visual({'type': 'UNKNOWN'}))
        self.assertIsNone(visuals.clean_visual({'type': 'SEQUENCE', 'items': ['solo uno']}))
        self.assertEqual(visuals.clean_visual({'type': 'FRACTION_BAR', 'parts': 4, 'shaded': 9, 'crossed': 7}),
                         {'type': 'FRACTION_BAR', 'parts': 4, 'shaded': 4, 'crossed': 4})

    def test_every_type_renders_a_png(self):
        from . import visuals
        specs = [{'type': 'FRACTION_CIRCLE', 'parts': 4, 'shaded': 3}, {'type': 'FRACTION_BAR', 'parts': 8, 'shaded': 5},
                 {'type': 'NUMBER_LINE', 'start': 0, 'end': 10, 'marks': [3]}, {'type': 'COUNT', 'count': 7, 'crossed': 2},
                 {'type': 'SEQUENCE', 'items': ['Lee', 'Responde']}]
        for spec in specs:
            self.assertTrue(visuals.render_png(visuals.clean_visual(spec)).startswith(b'\x89PNG'))
