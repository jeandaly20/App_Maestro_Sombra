"""
Original figures of the worksheet.

Every input becomes a list of page images that Gemini sees ("Imagen 1", "Imagen 2"...):
the photo itself, each page of a PDF, or the pictures embedded in a Word file. Gemini
returns where each figure is (box_2d, 0-1000 scale) and the system crops it, so the
adapted worksheet shows the student the real drawing of their guide, not a new one.
"""
import base64
import io
from pathlib import Path

import pypdfium2
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from docx import Document
from PIL import Image

from apps.materials.models import InputType

MAX_PDF_PAGES = 10
MAX_WORD_IMAGES = 6
MAX_SIDE = 1600          # px sent to the AI: enough to read small print
MIN_FIGURE_SIDE = 40     # px: smaller crops are noise


def _fit(image):
    image = image.convert('RGB')
    image.thumbnail((MAX_SIDE, MAX_SIDE))
    return image


def _jpeg(image):
    buffer = io.BytesIO()
    image.save(buffer, 'JPEG', quality=88)
    return buffer.getvalue()


def load_sources(material):
    """
    Returns (text, images): the text for Word/typed input and the list of PIL images
    the AI will look at, in order.
    """
    if material.input_type == InputType.TEXT:
        return material.input_text, []
    material.file.open('rb')
    try:
        raw = material.file.read()
    finally:
        material.file.close()
    if material.input_type == InputType.PHOTO:
        return '', [_fit(Image.open(io.BytesIO(raw)))]
    if material.input_type == InputType.PDF:
        pdf = pypdfium2.PdfDocument(raw)
        try:
            pages = [_fit(pdf[i].render(scale=2).to_pil()) for i in range(min(len(pdf), MAX_PDF_PAGES))]
        finally:
            pdf.close()
        return '', pages
    # Word: its text plus the pictures embedded in it
    document = Document(io.BytesIO(raw))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(' | '.join(cells))
    images = []
    for part in document.part.related_parts.values():
        if part.content_type.startswith('image/') and len(images) < MAX_WORD_IMAGES:
            try:
                image = Image.open(io.BytesIO(part.blob))
                if min(image.size) >= MIN_FIGURE_SIDE:
                    images.append(_fit(image))
            except OSError:
                continue  # formats Pillow cannot open (e.g. EMF) are skipped
    return '\n'.join(lines), images


def image_parts(images):
    """Gemini parts for the images, each one preceded by its number."""
    parts = []
    for number, image in enumerate(images, start=1):
        parts.append({'text': f'Imagen {number}:'})
        parts.append({'inlineData': {'mimeType': 'image/jpeg', 'data': base64.b64encode(_jpeg(image)).decode()}})
    return parts


def crop_figures(material, images, figures):
    """
    Crops each figure Gemini located and stores it (owner-only access through a view).
    Adds figure['image'] with the storage path; figures without a valid box keep only text.
    """
    folder = f'figures/{material.teacher_id}/{material.id}'
    for number, figure in enumerate(figures, start=1):
        source = figure.get('source') or 0
        box = figure.get('box_2d') or []
        if not (1 <= source <= len(images)) or len(box) != 4:
            continue
        image = images[source - 1]
        width, height = image.size
        ymin, xmin, ymax, xmax = [max(0, min(1000, int(v))) for v in box]
        if ymax <= ymin or xmax <= xmin:
            continue
        # Generous margin: the model's boxes are often a little tight and cut the edges of tables
        pad_x, pad_y = width * 0.03, height * 0.03
        crop_box = (max(0, xmin * width / 1000 - pad_x), max(0, ymin * height / 1000 - pad_y),
                    min(width, xmax * width / 1000 + pad_x), min(height, ymax * height / 1000 + pad_y))
        crop = image.crop(tuple(int(v) for v in crop_box))
        if min(crop.size) < MIN_FIGURE_SIDE:
            continue
        buffer = io.BytesIO()
        crop.save(buffer, 'PNG', optimize=True)
        path = default_storage.save(f'{folder}/figura_{number}.png', ContentFile(buffer.getvalue()))
        figure['image'] = path
    return figures


def delete_figures(material):
    """Removes the cropped figures of a material (used when it is deleted)."""
    for figure in (material.ai_analysis or {}).get('figures', []):
        path = figure.get('image')
        if path and default_storage.exists(path):
            default_storage.delete(path)


def figure_path_is_safe(path, material):
    """Only paths inside this material's folder can be served."""
    expected = Path('figures') / str(material.teacher_id) / str(material.id)
    return Path(path).parent == expected


def file_images(raw):
    """Page images of an uploaded PDF or photo (used for the teacher guide of a class)."""
    if raw[:5] == b'%PDF-':
        pdf = pypdfium2.PdfDocument(raw)
        try:
            return [_fit(pdf[i].render(scale=2).to_pil()) for i in range(min(len(pdf), MAX_PDF_PAGES))]
        finally:
            pdf.close()
    return [_fit(Image.open(io.BytesIO(raw)))]


def to_pdf(uploads):
    """Several photos (or one PDF) of the same pages as a single PDF, so they travel as one material."""
    if len(uploads) == 1 and uploads[0].read(5) == b'%PDF-':
        uploads[0].seek(0)
        return uploads[0].read()
    pages = []
    for upload in uploads:
        upload.seek(0)
        pages.append(_fit(Image.open(io.BytesIO(upload.read()))))
    buffer = io.BytesIO()
    pages[0].save(buffer, 'PDF', save_all=True, append_images=pages[1:], resolution=150)
    return buffer.getvalue()
