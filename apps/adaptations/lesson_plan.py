"""
The teacher's class plan (Excel) for "Adaptar clase".

Every school has its own format, so nothing here depends on fixed cells: the sessions are found by
their label ("Sesión 30", "Reto 15", "Clase 4") and each session becomes plain "label: value" text
that the AI interprets. Before that text leaves the server, the rows that identify the school,
the teachers or the student (institution, motto, names, signatures, dates, the current
"Adaptación curricular" row with the student's initials) are dropped.
"""
import io
import re
import unicodedata
from datetime import date, datetime

import openpyxl

SESSION_LABEL = re.compile(r'^\s*(sesi[oó]n|reto|clase)\s*n?[º°.]?\s*\d+', re.IGNORECASE)
# Rows with these words identify people or the school, or are administrative: never sent to the AI
PRIVATE_WORDS = ('docente', 'nombre', 'firma', 'revisado', 'aprobado', 'realizado', 'elaborado', 'unidad educativa',
                 'escuela', 'colegio', 'institucion', 'formato', 'actualiz', 'pagina', 'paralelo', 'fecha', 'periodo',
                 'ano lectivo', 'jefe de area', 'coordinacion', 'rector', 'observaciones', 'bibliografia',
                 'adaptacion curricular', 'estudiante', 'alumno:')
# A plan's pedagogical part starts at its first row with one of these labels (before it: banner, motto)
PEDAGOGICAL_WORDS = ('destreza', 'indicador', 'objetivo', 'eje', 'sesion', 'reto', 'tema', 'calentemos',
                     'inicio', 'anticipacion', 'desarrollo', 'contenido')
URL = re.compile(r'https?://\S+')
MAX_TEXT = 7000
ADAPTATION_LABEL = 'adaptacion curricular'


def plain(text):
    text = unicodedata.normalize('NFKD', str(text or ''))
    return ''.join(c for c in text if not unicodedata.combining(c)).lower().strip()


def open_workbook(content, *, values=True):
    """values=True reads computed values (for the text); False keeps formulas and styles (to write back)."""
    return openpyxl.load_workbook(io.BytesIO(content), data_only=values)


def _cell_text(value):
    if value is None or isinstance(value, (datetime, date)):
        return ''  # dates identify the school year and the class day: not needed to adapt
    return ' '.join(str(value).split())


def _rows(sheet):
    for row in sheet.iter_rows():
        texts = [_cell_text(cell.value) for cell in row]
        texts = [text for text in texts if text]
        if texts:
            yield row[0].row, texts


def sessions(content):
    """Every session found in the workbook: [{'key', 'sheet', 'row', 'label', 'title'}]."""
    workbook = open_workbook(content)
    found = []
    for sheet in workbook.worksheets:
        sheet_sessions = []
        for number, texts in _rows(sheet):
            label = next((text for text in texts if SESSION_LABEL.match(text)), None)
            if not label:
                continue
            # The title is usually the value after a "Título…" label in the same row
            title = ''
            for index, text in enumerate(texts):
                if plain(text).startswith('titulo') and index + 1 < len(texts):
                    title = texts[index + 1]
            sheet_sessions.append({'key': f'{sheet.title}|{number}', 'sheet': sheet.title, 'row': number,
                                   'label': label, 'title': title})
        if not sheet_sessions and sum(1 for _ in _rows(sheet)) > 5:
            # A plan without numbered sessions: the whole sheet is one class
            sheet_sessions.append({'key': f'{sheet.title}|0', 'sheet': sheet.title, 'row': 0,
                                   'label': 'Hoja completa', 'title': ''})
        found += sheet_sessions
    return found


def _is_private(texts):
    joined = plain(' '.join(texts))
    first = plain(texts[0])
    return any(first.startswith(word) or f' {word}' in f' {first}' for word in PRIVATE_WORDS) or (
        len(texts) == 1 and any(word in joined for word in ('unidad educativa', 'formando', 'lema')))


def session_text(content, sheet_name, row):
    """The chosen session as 'label: value' lines, without personal or school data."""
    sheet = open_workbook(content)[sheet_name]
    rows = list(_rows(sheet))
    starts = [number for number, texts in rows if any(SESSION_LABEL.match(text) for text in texts)]
    pedagogical = next((number for number, texts in rows
                        if any(word in plain(' '.join(texts)) for word in PEDAGOGICAL_WORDS)), 0)
    first_session = starts[0] if starts else None
    end = next((start for start in starts if start > row), None) if row else None
    lines = []
    for number, texts in rows:
        shared = pedagogical <= number < (first_session or 10 ** 9)  # sheet-level fields: skill, indicator…
        in_session = (not row) or (row <= number and (end is None or number < end))
        if not (shared or in_session) or number < pedagogical or _is_private(texts):
            continue
        lines.append(' | '.join(URL.sub('', text).strip() for text in texts if URL.sub('', text).strip()))
    return '\n'.join(line for line in lines if line)[:MAX_TEXT]


def write_adaptation(content, sheet_name, row, text):
    """
    The same workbook with the session's "Adaptación curricular" cell filled in, or None when the
    plan has no such row. Only that cell changes; the rest of the teacher's file stays as it was.
    """
    workbook = open_workbook(content, values=False)
    sheet = workbook[sheet_name]
    rows = [r for r in sheet.iter_rows() if any(cell.value not in (None, '') for cell in r)]
    starts = [r[0].row for r in rows if any(SESSION_LABEL.match(_cell_text(c.value)) for c in r)]
    end = next((start for start in starts if start > row), None) if row else None
    for cells in rows:
        number = cells[0].row
        if number < row or (end is not None and number >= end):
            continue
        for index, cell in enumerate(cells):
            if plain(_cell_text(cell.value)).startswith(ADAPTATION_LABEL):
                target = next((c for c in cells[index + 1:] if type(c).__name__ != 'MergedCell'), None)
                if target is None:
                    return None
                target.value = text
                target.alignment = target.alignment.copy(wrap_text=True, vertical='top')
                buffer = io.BytesIO()
                workbook.save(buffer)
                return buffer.getvalue()
    return None
